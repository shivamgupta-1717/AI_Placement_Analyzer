import io
import re
from typing import List, Dict, Any, Tuple, Optional
from dataclasses import dataclass, field
import pymupdf as fitz
from docx import Document

class DocumentParsingError(Exception):
    """Base exception for document parsing failures."""
    pass

class UnsupportedFileTypeError(DocumentParsingError):
    """Raised when file extension or MIME type is not supported."""
    pass

class FileSizeLimitExceededError(DocumentParsingError):
    """Raised when file size exceeds maximum permitted limit."""
    pass

class PasswordProtectedError(DocumentParsingError):
    """Raised when PDF or document is encrypted/password protected."""
    pass

class CorruptedDocumentError(DocumentParsingError):
    """Raised when document cannot be opened or parsed."""
    pass

class EmptyDocumentError(DocumentParsingError):
    """Raised when document contains no readable text."""
    pass

class ScannedDocumentError(DocumentParsingError):
    """Raised when document appears to be a scanned image with no text layer."""
    pass


@dataclass
class ParsedChunk:
    chunk_index: int
    content: str
    section_name: Optional[str]
    page_number: Optional[int]
    token_count: int


@dataclass
class DocumentParseResult:
    filename: str
    file_type: str
    file_size: int
    full_text: str
    detected_sections: List[str]
    chunks: List[ParsedChunk]
    metadata: Dict[str, Any] = field(default_factory=dict)


# Standard resume section header patterns
SECTION_PATTERNS = [
    (r"(?i)^(?:technical\s+)?skills(?:\s+summary|\s+&\s+proficiencies)?(?:\s*:)?$", "Skills"),
    (r"(?i)^(?:work\s+)?experience(?:\s+history|\s+summary)?(?:\s*:)?$", "Experience"),
    (r"(?i)^(?:employment\s+history|internships?)(?:\s*:)?$", "Experience"),
    (r"(?i)^(?:projects|academic\s+projects|personal\s+projects)(?:\s*:)?$", "Projects"),
    (r"(?i)^(?:education|academic\s+background|qualifications)(?:\s*:)?$", "Education"),
    (r"(?i)^(?:certifications?|licenses?|credentials?)(?:\s*:)?$", "Certifications"),
    (r"(?i)^(?:achievements|awards|honors)(?:\s*:)?$", "Achievements"),
    (r"(?i)^(?:summary|professional\s+summary|profile|about\s+me)(?:\s*:)?$", "Summary"),
    (r"(?i)^(?:publications|patents)(?:\s*:)?$", "Publications"),
]

def normalize_text(text: str) -> str:
    """Normalizes whitespace while preserving paragraphs and line boundaries."""
    if not text:
        return ""
    # Normalize unicode spaces and non-breaking spaces
    text = text.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    # Replace sequences of 3+ newlines with 2 newlines
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Replace non-newline whitespace sequences with a single space
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(lines).strip()

def detect_section_header(line: str) -> Optional[str]:
    """Detects if a single line acts as a major resume section heading."""
    clean_line = line.strip(" :-\t#*")
    if not clean_line or len(clean_line) > 50:
        return None
    for pattern, normalized_name in SECTION_PATTERNS:
        if re.match(pattern, clean_line):
            return normalized_name
    return None

class DocumentParserService:
    def __init__(self, max_file_size_bytes: int = 10 * 1024 * 1024):
        self.max_file_size_bytes = max_file_size_bytes

    def parse_document(self, filename: str, content_bytes: bytes) -> DocumentParseResult:
        """
        Parses PDF or DOCX resume document, validates magic bytes, extracts text,
        detects sections, and generates semantic chunks.
        """
        if not content_bytes:
            raise EmptyDocumentError("The uploaded file is empty (0 bytes).")
        
        file_size = len(content_bytes)
        if file_size > self.max_file_size_bytes:
            raise FileSizeLimitExceededError(
                f"File size ({file_size / (1024*1024):.1f}MB) exceeds the maximum limit of {self.max_file_size_bytes / (1024*1024):.0f}MB."
            )

        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            return self._parse_pdf(filename, content_bytes)
        elif lower_name.endswith(".docx"):
            return self._parse_docx(filename, content_bytes)
        else:
            raise UnsupportedFileTypeError(
                f"Unsupported file format '{filename}'. Please upload a .pdf or .docx resume."
            )

    def _parse_pdf(self, filename: str, content_bytes: bytes) -> DocumentParseResult:
        # Check PDF magic bytes '%PDF-'
        if not content_bytes.startswith(b"%PDF-"):
            raise CorruptedDocumentError("Invalid PDF format: Missing standard PDF header magic bytes.")

        try:
            doc = fitz.open(stream=content_bytes, filetype="pdf")
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to read PDF file: {str(e)}")

        if doc.is_encrypted:
            doc.close()
            raise PasswordProtectedError("The uploaded PDF is password protected. Please upload an unprotected copy.")

        total_pages = len(doc)
        if total_pages == 0:
            doc.close()
            raise EmptyDocumentError("The uploaded PDF has no pages.")

        page_texts: List[Tuple[int, str]] = []
        total_chars = 0

        for page_num in range(total_pages):
            page = doc[page_num]
            text = page.get_text("text") or ""
            cleaned = normalize_text(text)
            total_chars += len(cleaned)
            page_texts.append((page_num + 1, cleaned))

        doc.close()

        # Check if scanned / image-only PDF
        if total_chars < 50:
            raise ScannedDocumentError(
                "No readable text found in PDF. This appears to be a scanned image-only document. "
                "Please upload a searchable text-based PDF or DOCX file (or perform OCR before uploading)."
            )

        full_text = "\n\n".join([t for _, t in page_texts if t])
        chunks, detected_sections = self._chunk_document_pages(page_texts)

        return DocumentParseResult(
            filename=filename,
            file_type="application/pdf",
            file_size=len(content_bytes),
            full_text=full_text,
            detected_sections=detected_sections,
            chunks=chunks,
            metadata={"total_pages": total_pages, "total_characters": total_chars}
        )

    def _parse_docx(self, filename: str, content_bytes: bytes) -> DocumentParseResult:
        # Check DOCX zip magic bytes 'PK\x03\x04'
        if not content_bytes.startswith(b"PK\x03\x04"):
            raise CorruptedDocumentError("Invalid DOCX format: Missing standard DOCX archive signature.")

        try:
            doc = Document(io.BytesIO(content_bytes))
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to read DOCX file: {str(e)}")

        paragraphs = []
        for p in doc.paragraphs:
            text = normalize_text(p.text)
            if text:
                paragraphs.append(text)

        # Also extract table text safely
        for table in doc.tables:
            for row in table.rows:
                row_cells = [normalize_text(cell.text) for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))

        if not paragraphs:
            raise EmptyDocumentError("The uploaded DOCX file contains no readable text content.")

        full_text = "\n\n".join(paragraphs)
        if len(full_text) < 50:
            raise EmptyDocumentError("The uploaded DOCX document contains insufficient text content.")

        # Single page equivalent for docx
        page_texts = [(1, full_text)]
        chunks, detected_sections = self._chunk_document_pages(page_texts)

        return DocumentParseResult(
            filename=filename,
            file_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_size=len(content_bytes),
            full_text=full_text,
            detected_sections=detected_sections,
            chunks=chunks,
            metadata={"paragraph_count": len(paragraphs), "total_characters": len(full_text)}
        )

    def _chunk_document_pages(self, page_texts: List[Tuple[int, str]]) -> Tuple[List[ParsedChunk], List[str]]:
        """
        Splits document pages into semantic chunks while tracking section context and page numbers.
        """
        chunks: List[ParsedChunk] = []
        detected_sections_set = set()
        current_section = "General"
        chunk_idx = 0

        for page_num, text in page_texts:
            lines = text.split("\n")
            current_chunk_lines: List[str] = []
            current_chunk_length = 0

            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue

                # Check if this line is a section header
                new_sec = detect_section_header(line_str)
                if new_sec:
                    # Flush current chunk before starting new section
                    if current_chunk_lines:
                        chunk_content = "\n".join(current_chunk_lines).strip()
                        if chunk_content:
                            chunks.append(ParsedChunk(
                                chunk_index=chunk_idx,
                                content=chunk_content,
                                section_name=current_section,
                                page_number=page_num,
                                token_count=max(1, len(chunk_content.split()))
                            ))
                            chunk_idx += 1
                        current_chunk_lines = []
                        current_chunk_length = 0
                    current_section = new_sec
                    detected_sections_set.add(new_sec)
                    current_chunk_lines.append(f"[{new_sec}]")
                    continue

                # Add line to current chunk
                current_chunk_lines.append(line_str)
                current_chunk_length += len(line_str)

                # Chunk size target: ~400-600 characters or logical paragraph boundary
                if current_chunk_length >= 450:
                    chunk_content = "\n".join(current_chunk_lines).strip()
                    if chunk_content:
                        chunks.append(ParsedChunk(
                            chunk_index=chunk_idx,
                            content=chunk_content,
                            section_name=current_section,
                            page_number=page_num,
                            token_count=max(1, len(chunk_content.split()))
                        ))
                        chunk_idx += 1
                    current_chunk_lines = []
                    current_chunk_length = 0

            # Flush any remaining lines on this page
            if current_chunk_lines:
                chunk_content = "\n".join(current_chunk_lines).strip()
                if chunk_content:
                    chunks.append(ParsedChunk(
                        chunk_index=chunk_idx,
                        content=chunk_content,
                        section_name=current_section,
                        page_number=page_num,
                        token_count=max(1, len(chunk_content.split()))
                    ))
                    chunk_idx += 1

        detected_sections = sorted(list(detected_sections_set))
        return chunks, detected_sections
