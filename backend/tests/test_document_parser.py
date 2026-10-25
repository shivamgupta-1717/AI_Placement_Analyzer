import pytest
import pymupdf as fitz
import io
from docx import Document
from app.services.document_parser import (
    DocumentParserService,
    UnsupportedFileTypeError,
    FileSizeLimitExceededError,
    PasswordProtectedError,
    CorruptedDocumentError,
    EmptyDocumentError,
    ScannedDocumentError
)

@pytest.fixture
def parser():
    return DocumentParserService()

def test_unsupported_file_extension(parser):
    with pytest.raises(UnsupportedFileTypeError):
        parser.parse_document("resume.exe", b"fake binary")

def test_empty_file_error(parser):
    with pytest.raises(EmptyDocumentError):
        parser.parse_document("resume.pdf", b"")

def test_corrupted_pdf(parser):
    with pytest.raises(CorruptedDocumentError):
        parser.parse_document("resume.pdf", b"not a real pdf content")

def test_valid_pdf_parsing(parser):
    # Create a real in-memory PDF with PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "SKILLS\nPython, Docker, React, PostgreSQL\n\nEXPERIENCE\nSoftware Engineer at Tech Corp\nBuilt REST APIs.")
    pdf_bytes = doc.tobytes()
    doc.close()

    result = parser.parse_document("test_resume.pdf", pdf_bytes)
    assert result.filename == "test_resume.pdf"
    assert "Python" in result.full_text
    assert len(result.chunks) > 0
    assert "Skills" in result.detected_sections or "Experience" in result.detected_sections

def test_scanned_pdf_detection(parser):
    # PDF with almost zero characters
    doc = fitz.open()
    doc.new_page()
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(ScannedDocumentError):
        parser.parse_document("scanned.pdf", pdf_bytes)

def test_valid_docx_parsing(parser):
    doc = Document()
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Java, Spring Boot, Microservices, Kubernetes.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Technology in Computer Science.")
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = parser.parse_document("test_resume.docx", docx_bytes)
    assert result.filename == "test_resume.docx"
    assert "Spring Boot" in result.full_text
    assert len(result.chunks) > 0
